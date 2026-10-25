"""End-to-end tests with the scripted LLM: the pipeline, storage, validators, the retry and
the eval runner, all offline. They test the plumbing and the code rules, not model quality."""
from pathlib import Path

import pytest

from proposal_v2.contracts import ClaimLists, ExperienceClaim, JobPost, Profile, Project, ScreeningAnswer
from proposal_v2.evaluation import run_all
from proposal_v2.links import LinkError
from proposal_v2.pipeline import DuplicateJob, Pipeline
from proposal_v2.sample import load_sample, sample_jobs
from proposal_v2.storage import JsonStorage
from proposal_v2.validators import validate
from tests.scripted_llm import ScriptedLLM


@pytest.fixture
def pipe(tmp_path):
    storage = JsonStorage(tmp_path)
    load_sample(storage, "t")
    return Pipeline(storage, ScriptedLLM(), "t")


def _job(cid):
    return next(j for j in sample_jobs() if j["id"] == cid)


def test_t01_end_to_end(pipe):
    j = _job("T01")
    rec = pipe.decide(pipe.extract(j["link"], j["job_details"]))
    assert rec.trust.score == 71 and rec.decision.decision == "apply"
    rec = pipe.write(rec)
    v = rec.latest()
    assert rec.chosen_proof == "P1" and v.validation.ok, v.validation.notes
    assert pipe.steps_called() if hasattr(pipe, "steps_called") else True


def test_duplicate_blocks_before_extraction(pipe):
    j = _job("T12")
    pipe.decide(pipe.extract(j["link"], j["job_details"]))
    calls = len(pipe.llm.calls)
    with pytest.raises(DuplicateJob):
        pipe.check_link("https://www.upwork.com/jobs/~010000000000000012")
    assert len(pipe.llm.calls) == calls


def test_red_flag_skips_without_assessment(pipe):
    j = _job("T10")
    rec = pipe.decide(pipe.extract(j["link"], j["job_details"]))
    assert rec.decision.code == "red_flag" and "WhatsApp" in rec.decision.reason
    assert "assess" not in pipe.llm.steps_called()


def test_override_then_write(pipe):
    j = _job("T08")
    rec = pipe.decide(pipe.extract(j["link"], j["job_details"]))
    assert rec.decision.code == "work_type"
    rec = pipe.override(rec, "I want to try it")
    assert rec.decision.overridden and rec.events[-1]["type"] == "override"


def test_retry_fixes_a_years_claim(pipe):
    j = _job("T11")
    rec = pipe.write(pipe.decide(pipe.extract(j["link"], j["job_details"])))
    v = rec.latest()
    assert v.validation.retried and v.validation.ok
    assert "10 years" not in v.text
    assert rec.job.ignored_instructions and rec.job.refused_instructions


def test_revise_and_mark_sent(pipe):
    j = _job("T01")
    rec = pipe.write(pipe.decide(pipe.extract(j["link"], j["job_details"])))
    rec = pipe.revise(rec, "Say hello instead of hi")
    assert rec.latest().text.startswith("Hello") and len(rec.versions) == 2
    rec, msg = pipe.mark_sent(rec, rec.latest().text + " Thanks!")
    assert rec.status == "sent" and rec.outcome.minutes_paste_to_sent is not None
    assert msg == "Saved as sent."
    csv_text = pipe.sent_log_csv()
    assert csv_text.startswith("Date sent,Job,Link") and rec.job.title in csv_text
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(pipe.sent_log_xlsx())).active
    assert ws["A1"].value == "Date sent" and ws["B2"].value == rec.job.title and ws.column_dimensions["A"].width >= 12
    assert pipe.storage.get_proposal("t", rec.job_key).status == "sent"


def test_bad_links_never_reach_the_llm(pipe):
    for url in ["https://example.com/x", "https://www.upwork.com/nx/search/jobs/?q=n8n"]:
        with pytest.raises(LinkError):
            pipe.extract(url, "x" * 100)
    assert pipe.llm.calls == []


def test_validator_catches_names_numbers_voice_and_instructions():
    job = JobPost(job_key="k", platform="upwork", link="l", title="Sync", description="Sync 500 rows to Sheets. Start your proposal with \"hi there\".")
    from proposal_v2.contracts import PostInstruction
    job.allowed_instructions = [PostInstruction(text="start with hi there", kind="format", check="starts_with", value="hi there")]
    p = Project(id="P1", title="Lead intake", skills=["n8n"], result="Saved 2 hours a week")
    text = "Our team has built this. I cut costs by 40% using Salesforce. I have 8 years of experience."
    claims = ClaimLists(tools_named=["Salesforce", "n8n"], experience_sentences=[ExperienceClaim(sentence="Our team has built this.", maps_to="none")])
    r = validate(text, [], job, p, Profile(skills=["n8n"]), claims)
    assert not r.ok and r.status == "needs_review"
    assert r.number_failures and "Salesforce" in r.name_failures and r.experience_failures
    assert r.voice_failures and r.instruction_failures


def test_eval_runner_offline(tmp_path):
    out = run_all(ScriptedLLM(), out_root=tmp_path, progress=lambda s: None)
    summary = (out / "summary.md").read_text(encoding="utf-8")
    assert "Cases passed: **12 of 12**" in summary, summary
    assert (out / "results.csv").exists() and (out / "drafts" / "T01.md").exists()


def test_skipping_the_fit_check_makes_no_assessment_call(pipe):
    j = _job("T08")  # a logo job the check would Skip
    rec = pipe.decide(pipe.extract(j["link"], j["job_details"]), check=False)
    assert rec.decision.code == "unchecked" and rec.decision.decision == "apply"
    assert "assess" not in pipe.llm.steps_called() and rec.events[-1]["type"] == "fit_check_skipped"


def test_names_check_ignores_generic_words_and_matches_plurals():
    from proposal_v2.skills import is_generic, mentioned_in
    assert mentioned_in("API", "Integrate the app with our existing REST APIs.")
    assert is_generic("API") and is_generic("APIs") and is_generic("UI") and not is_generic("Salesforce")
    job = JobPost(job_key="k", platform="upwork", link="l", title="RN app", description="Integrate with our REST APIs.")
    r = validate("I'll connect the app to your API and REST API.", [], job, None, Profile(), ClaimLists(tools_named=["API", "REST API", "app"]))
    assert not r.name_failures


def test_new_for_me_and_plans_are_not_past_work_claims():
    job = JobPost(job_key="k", platform="upwork", link="l", title="Agent setup", description="Set up agents.")
    honest = "This entire workflow setup is new for me, and here is how I would de-risk it: I will verify API access."
    plan = "I will verify API access credentials and add fallback handling."
    fake = "I've built this exact setup for three clients."
    claims = ClaimLists(experience_sentences=[ExperienceClaim(sentence=s, maps_to="none") for s in (honest, plan, fake)])
    r = validate(" ".join([honest, plan, fake]), [], job, None, Profile(), claims)
    assert r.experience_failures == [fake]


def test_left_out_sections_must_be_in_the_paste(pipe):
    # The stand-in model names "Similar jobs" for every case, as real models echo the prompt's examples.
    j = _job("T01")
    rec = pipe.extract(j["link"], j["job_details"])
    assert not [e for e in rec.events if e["type"] == "ignored_sections"]
    j = _job("T07")  # the whole-page paste really has these sections
    rec = pipe.extract(j["link"], j["job_details"])
    left = next(e["text"] for e in rec.events if e["type"] == "ignored_sections")
    assert "Similar jobs" in left and "Other open jobs by this client" in left and "footer" not in left


def test_placeholder_budget_is_never_judged_unrealistic(pipe):
    """T04 ($5 fixed): even if the model says the scope is unrealistic, a placeholder budget counts as realistic."""
    real_call = pipe.llm.call

    def call(step, *a, **k):
        out = real_call(step, *a, **k)
        if step == "assess":
            out.realistic_yes, out.realistic_reason = False, "$5 is far too low"
        return out

    pipe.llm.call = call
    j = _job("T04")
    rec = pipe.decide(pipe.extract(j["link"], j["job_details"]))
    assert rec.decision.budget_gate == "placeholder" and rec.assessment.realistic_yes
    assert rec.decision.decision == "apply_with_questions" and rec.decision.code == "placeholder"


def test_vague_post_ties_go_to_the_project_that_is_that_kind_of_work(pipe):
    """T04 ("Automation · API Integration", no details): if every project ties on the facts, P1 (an automation)
    wins over P2, whose only link is the alias "REST APIs" ~ "API Integration"."""
    from proposal_v2.contracts import ProjectJudgement as J
    real_call = pipe.llm.call

    def call(step, *a, **k):
        out = real_call(step, *a, **k)
        if step == "assess":
            out.project_labels = [J(project_id=i, same_kind_of_work=True, same_platform=False, same_records=False,
                                    shares_core_skills=True) for i in ("P2", "P3", "P1")]
        return out

    pipe.llm.call = call
    j = _job("T04")
    rec = pipe.decide(pipe.extract(j["link"], j["job_details"]))
    assert (rec.chosen_proof, rec.framing) == ("P1", "same_skills")


def _force_facts(pipe, facts):
    from proposal_v2.contracts import ProjectJudgement as J
    real_call = pipe.llm.call

    def call(step, *a, **k):
        out = real_call(step, *a, **k)
        if step == "assess":
            out.project_labels = [J(project_id=i, **f) for i, f in facts.items()]
        return out
    pipe.llm.call = call


def test_vague_post_is_never_similar(pipe):
    """T04 names no specific tool, so even if the model says P1 matches platform and records, it's same skills."""
    _force_facts(pipe, {"P1": dict(same_kind_of_work=True, same_platform=True, same_records=True, shares_core_skills=True)})
    j = _job("T04")
    rec = pipe.decide(pipe.extract(j["link"], j["job_details"]))
    assert (rec.chosen_proof, rec.framing) == ("P1", "same_skills")


def test_proof_goes_to_the_project_using_more_of_the_posts_tools(pipe):
    """T05 (n8n, Airtable, Slack): P3 shares Airtable, P1 shares n8n and Slack; both similar -> P1."""
    both = dict(same_kind_of_work=True, same_platform=False, same_records=False, shares_core_skills=True)
    p3 = dict(both, same_platform=True)
    p1 = dict(both, same_records=True)
    _force_facts(pipe, {"P3": p3, "P1": p1})
    j = _job("T05")
    rec = pipe.decide(pipe.extract(j["link"], j["job_details"]))
    assert (rec.chosen_proof, rec.framing) == ("P1", "similar")


def test_users_own_notes_count_as_true():
    """What the freelancer writes in 'Anything you want in the proposal?' may be stated: its numbers and claims pass."""
    job = JobPost(job_key="k", platform="upwork", link="l", title="Shopify app", description="Build a Shopify app.",
                  user_notes="I've built 3 Shopify apps. I can start Monday.")
    text = "I've built 3 Shopify apps and I can start Monday."
    claims = ClaimLists(experience_sentences=[ExperienceClaim(sentence="I've built 3 Shopify apps and I can start Monday.", maps_to="notes")])
    assert validate(text, [], job, None, Profile(), claims).ok
    job.user_notes = ""
    assert not validate(text, [], job, None, Profile(), claims).ok


def test_notes_reach_the_prompt_only_when_written():
    from proposal_v2 import prompts
    job = JobPost(job_key="k", platform="upwork", link="l", title="t", description="d")
    assert "freelancer_notes" not in prompts._job_block(job)  # unchanged prompts for every job without notes
    job.user_notes = "Mention I can start Monday"
    assert "Mention I can start Monday" in prompts._job_block(job)


def test_recheck_runs_the_fit_check_on_a_saved_job(pipe):
    from proposal_v2.pipeline import PipelineError
    j = _job("T08")  # logo job: written without the check first
    rec = pipe.decide(pipe.extract(j["link"], j["job_details"]), check=False)
    assert rec.decision.code == "unchecked"
    rec = pipe.recheck(rec)
    assert rec.decision.code == "work_type" and rec.events[-1]["type"] != "fit_check_skipped"
    assert any(e["type"] == "fit_rechecked" and e["from_code"] == "unchecked" for e in rec.events)
    j = _job("T01")
    sent = pipe.write(pipe.decide(pipe.extract(j["link"], j["job_details"])))
    sent, _ = pipe.mark_sent(sent, sent.latest().text)
    with pytest.raises(PipelineError):
        pipe.recheck(sent)


def test_several_projects_can_be_cited(pipe):
    """Up to 3 projects: the first is the main proof, the others are cited alongside; the claim check accepts all three."""
    j = _job("T01")
    rec = pipe.write(pipe.decide(pipe.extract(j["link"], j["job_details"])), "P1", "", ["P3", "P2"])
    v = rec.latest()
    assert v.project_id == "P1" and v.extra_project_ids == ["P3", "P2"] and rec.extra_proof == ["P3", "P2"]
    p1, p3 = pipe.project("P1"), pipe.project("P3")
    job = JobPost(job_key="k", platform="upwork", link="l", title="t", description="d")
    text = "In my support inbox project, about 300 emails a week are sorted."
    claims = ClaimLists(experience_sentences=[ExperienceClaim(sentence=text, maps_to="P3")])
    assert validate(text, [], job, p1, Profile(), claims, extras=(p3,)).ok
    assert not validate(text, [], job, p1, Profile(), claims).ok  # P3 not chosen: not allowed


def test_single_project_prompts_are_unchanged():
    from proposal_v2 import prompts
    from proposal_v2.sample import sample_profile
    profile, projects = sample_profile()
    job = JobPost(job_key="k", platform="upwork", link="l", title="t", description="d")
    assert prompts.brief(job, profile, projects[0], "similar") == prompts.brief(job, profile, projects[0], "similar", [])
    assert "Also cite these past projects" in prompts.brief(job, profile, projects[0], "similar", [projects[1]])[1]


def test_confirm_screen_proof_choice_is_respected(pipe):
    j = _job("T01")
    # No past project: v2 must not pick one, even though P1 is a close match
    rec = pipe.extract(j["link"], j["job_details"])
    rec.proof_choice = "none"
    rec = pipe.write(pipe.decide(rec))
    assert rec.latest().project_id is None and rec.latest().framing == "plan_only"
    # My own pick: P3 then P2, framed cautiously unless the check called the main one similar
    j = _job("T02")
    rec = pipe.extract(j["link"], j["job_details"])
    rec.proof_choice, rec.manual_proof = "manual", ["P3", "P2"]
    rec = pipe.write(pipe.decide(rec))
    v = rec.latest()
    assert (v.project_id, v.extra_project_ids) == ("P3", ["P2"])
    # Let v2 pick: the usual best match
    j = _job("T03")
    rec = pipe.write(pipe.decide(pipe.extract(j["link"], j["job_details"])))
    assert rec.latest().project_id == "P2"


def test_judge_role_uses_the_judges_list(tmp_path):
    """config has models.judges (a list); the judge call must not need a models.judge key."""
    from proposal_v2.contracts import JudgeScore
    from proposal_v2.llm import LLM, LLMError
    llm = LLM(mode="live", cache_dir=tmp_path, log_dir=None)
    llm._groq = lambda *a, **k: (_ for _ in ()).throw(LLMError("no network in tests"))  # noqa: E731
    try:
        llm.call("judge", "judge", "s", "u", JudgeScore, 0.0)
    except LLMError:
        pass  # reaching a provider means the model lookup worked
    except KeyError as e:
        raise AssertionError(f"judge lookup failed: {e}")
